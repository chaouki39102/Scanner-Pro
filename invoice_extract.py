# -*- coding: utf-8 -*-
"""
invoice_extract.py — استخراج حقول الفواتير (عربي / فرنسي / إنجليزي) من نص OCR.

المبدأ: لا نثق بنتيجة OCR أبدًا. كل حقل يخرج مع حالته (ok / warn / missing / computed)
والسطر الذي أُخذ منه، ثم تُجرى فحوص حسابية (مجموع HT + TVA = TTC ...) لكشف أخطاء القراءة.
المراجعة البشرية تبقى هي الحَكَم؛ الوحدة تُسرّعها وتُظهر مكان الشك.
"""
import csv
import datetime as _dt
import itertools
import json
import re

# ----------------------------------------------------------------- تطبيع النص
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "0123456789" * 2)
_INVISIBLE = re.compile("[\u200e\u200f\u202a-\u202e\u2066-\u2069\ufeff\u200b\u200c\u200d]")


def normalize_text(text):
    """أرقام عربية/فارسية → لاتينية، إزالة علامات الاتجاه، توحيد الفواصل والمسافات."""
    t = (text or "").translate(_DIGITS)
    t = t.replace("٫", ".").replace("٬", ",").replace("،", ",")
    t = _INVISIBLE.sub("", t)
    t = t.replace("\u00a0", " ").replace("\u202f", " ").replace("\u2009", " ")
    t = t.replace("–", "-").replace("—", "-").replace("‑", "-")
    t = re.sub(r"[ \t]+", " ", t)
    return "\n".join(line.strip() for line in t.split("\n"))


# -------------------------------------------------------------------- المبالغ
_PCT = re.compile(r"\d{1,3}(?:[.,]\d+)?\s*%|%\s*\d{1,3}(?:[.,]\d+)?")
_NUM_TOKEN = re.compile(r"(?<![\w/.,])[-+]?\d[\d.,]*(?![\w/%])")


def parse_amount(tok):
    """يحوّل '12 345,60' أو '12.345,60' أو '12,345.60' إلى float. يعيد None إذا لم يكن مبلغًا."""
    s = re.sub(r"\s+", "", str(tok)).strip(".,")
    if not s or not re.fullmatch(r"[-+]?\d[\d.,]*", s):
        return None
    neg = s.startswith("-")
    s = s.lstrip("+-")
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        thou = "." if dec == "," else ","
        s = s.replace(thou, "").replace(dec, ".")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        parts = s.split(sep)
        if len(parts) > 2:                       # 1.234.567 → آلاف
            s = "".join(parts)
        elif len(parts[1]) == 3 and 1 <= len(parts[0]) <= 3:   # 1.234 → ألف (الفواتير تستعمل خانتين عشريتين)
            s = "".join(parts)
        else:
            s = parts[0] + "." + parts[1]
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if neg else v


def _amounts_in(line):
    """كل المبالغ في سطر مع مواضعها، بعد حذف النسب المئوية. يدمج مجموعات الآلاف المفصولة بمسافة."""
    clean = _PCT.sub(lambda m: " " * len(m.group(0)), line)
    out = []
    for m in re.finditer(r"(?<![\w/.,])[-+]?\d{1,3}(?:[ .,]\d{3})+(?:[.,]\d{1,2})?(?![\w/%\d])|"
                         r"(?<![\w/.,])[-+]?\d+(?:[.,]\d{1,3})?(?![\w/%\d])", clean):
        v = parse_amount(m.group(0))
        if v is not None:
            out.append((m.start(), v, m.group(0)))
    return out



# ----------------------------------------------- تصحيح ترتيب العربية المعكوس
_AR = re.compile(r"[\u0621-\u064A\u0671-\u06D3]")
_LEX = ("فاتورة رقم بتاريخ تاريخ السيد السادة الزبون العميل مؤسسة شركة المجموع مجموع خارج الرسم رسم القيمة المضافة "
        "مع على حقوق الطابع صافي للدفع الدفع التجاري السجل الجبائي الرقم التعريف الإحصائي الاحصائي الجبائية دج دينار "
        "الكمية السعر البيان الوحدة المبلغ الإجمالي الاجمالي نقدا شيك تحويل الاستحقاق طريقة الصناعية المنطقة الإجمالي "
        "بدون قبل المستحق الواجب دفعه المشتري").split()
_LEXSET = set(_LEX)
_REVSET = {w[::-1] for w in _LEX}


def _ar_tokens_hits(tokens, lexset):
    n = 0
    for t in tokens:
        w = t.strip(".,:;()[]|-_/\\")
        if w in lexset:
            n += 1
    return n


def _logical_from_visual(line):
    """سطر مقروء بترتيب بصري (يسار→يمين) → ترتيب منطقي: نعكس تتابع مقاطع العربية وحروف كل كلمة، ونُبقي مقاطع اللاتيني/الأرقام كما هي."""
    toks = line.split(" ")
    runs, cur, cur_ar = [], [], None
    for t in toks:
        is_ar = bool(_AR.search(t))
        if cur and is_ar != cur_ar:
            runs.append((cur_ar, cur))
            cur = []
        cur_ar = is_ar
        cur.append(t)
    if cur:
        runs.append((cur_ar, cur))
    out = []
    for is_ar, ts in reversed(runs):
        out.extend([t[::-1] for t in reversed(ts)] if is_ar else ts)
    return " ".join(out)


def fix_arabic_order(text):
    """يكشف الأسطر العربية المعكوسة (Tesseract يُخرجها أحيانًا بترتيب بصري) ويعيدها لترتيبها المنطقي.
    القرار بالقاموس: إن طابقت كلمات الفواتير بعد العكس أكثر مما قبله نعكس. الأسطر غير الحاسمة تتبع أغلبية المستند."""
    lines = text.split("\n")
    decisions = []
    for ln in lines:
        if len(_AR.findall(ln)) < 2:
            decisions.append(None)
            continue
        a = _ar_tokens_hits(ln.split(), _LEXSET)
        b = _ar_tokens_hits(ln.split(), _REVSET)
        decisions.append(True if b > a else False if a > b else "?")
    rev = sum(1 for d in decisions if d is True)
    log = sum(1 for d in decisions if d is False)
    default = rev > log and rev >= 2
    out = []
    for ln, d in zip(lines, decisions):
        if d is True or (d == "?" and default):
            out.append(_logical_from_visual(ln))
        else:
            out.append(ln)
    return "\n".join(out)



def heal_arabic(text):
    """يصلح كلمات القاموس المشوّهة قليلًا (ذارج→خارج) بتشابه نصي ≥ 0.74، على المستندات العربية فقط."""
    import difflib
    out = []
    for ln in text.split("\n"):
        toks = []
        for t in ln.split(" "):
            core = t.strip(".,:;()[]|-_/\\")
            if len(core) >= 4 and _AR.search(core) and core not in _LEXSET:
                m = difflib.get_close_matches(core, _LEX, n=1, cutoff=0.74)
                if m:
                    t = t.replace(core, m[0])
            toks.append(t)
        out.append(" ".join(toks))
    return "\n".join(out)


def _ids_by_shape(lines, limit_idx):
    """احتياط عند فشل قراءة تسميات المعرّفات: التصنيف بشكل الرقم (15 خانة: NIF ثم NIS، 11: AI، نمط RC)."""
    res = {}
    seq15 = []
    for ln in lines[: (limit_idx if limit_idx else 24)]:
        for m in re.finditer(r"(?<![\d,.])\d(?:[ ]?\d){14}(?![\d,])", ln):
            if len(re.sub(r"\D", "", m.group(0))) == 15:
                seq15.append((re.sub(r"\D", "", m.group(0)), ln))
        m = re.search(r"(?<![\d,.])\d(?:[ ]?\d){10}(?![\d,])", ln)
        if m and len(re.sub(r"\D", "", m.group(0))) == 11 and "ai" not in res and not re.search(r"\d{12,}", ln.replace(" ", "")):
            res["ai"] = (re.sub(r"\D", "", m.group(0)), ln)
        mr = re.search(r"(?<!\d)(\d{2})\s+(\d{6,8})\s*-\s*(\d{2})\s*/\s*(\d{2})(?!\d)", ln)
        if mr and "rc" not in res:
            res["rc"] = (f"{mr.group(3)}/{mr.group(4)}-{mr.group(2)} B {mr.group(1)}", ln)
        m = re.search(r"\d{1,2}\s*/\s*\d{2}\s*-\s*\d{6,8}(?:\s*[A-Za-z\u0621-\u064A0-9]\s*\d{2})?", ln)
        if m and "rc" not in res:
            res["rc"] = (re.sub(r"\s*([/\-])\s*", r"\1", m.group(0)), ln)
    if seq15:
        res["nif"] = seq15[0]
    if len(seq15) > 1:
        res["nis"] = seq15[1]
    return res


def _fix_rc(rc):
    """RC بصيغة NN/NN-NNNNNNN X NN: الحرف الأوسط غالبًا B أو A، و OCR يقرأ B كـ 8 كثيرًا."""
    m = re.fullmatch(r"(\d{1,2}/\d{2}-\d{6,8})\s+([A-Za-z0-9])\s+(\d{2})", rc or "")
    if m and m.group(2) in "84":
        return f"{m.group(1)} {'B' if m.group(2) == '8' else 'A'} {m.group(3)}", True
    return rc, False


# ------------------------------------------------------------------- التواريخ
_DATE_DMY = re.compile(r"(?<![\d/.\-])(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{4}|\d{2})(?![\d/])")
_DATE_YMD = re.compile(r"(?<![\d/.\-])(\d{4})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})(?![\d/])")
_DUE_WORDS = re.compile(r"[ée]ch[ée]ance|due|تاريخ\s*الاستحقاق|الاستحقاق|[àa]\s*payer\s*avant", re.I)
_DATE_WORDS = re.compile(r"\bdate\b|\ble\b|fait\s*[àa]|alger|تاريخ|بتاريخ|حرر\s*في|في\s*:", re.I)


def _mk_date(d, m, y):
    y = int(y)
    if y < 100:
        y += 2000 if y < 70 else 1900
    try:
        return _dt.date(y, int(m), int(d))
    except ValueError:
        return None


def _dates_in(line):
    res = []
    for m in _DATE_YMD.finditer(line):
        d = _mk_date(m.group(3), m.group(2), m.group(1))
        if d:
            res.append(d)
    for m in _DATE_DMY.finditer(line):
        d = _mk_date(m.group(1), m.group(2), m.group(3))
        if d:
            res.append(d)
    return res


# --------------------------------------------------------------- تسميات الحقول
_CODE = r"([A-Za-z0-9][A-Za-z0-9/\-_.]{1,26})"
_INV_PATTERNS = [
    re.compile(r"(?i)\bfacture(?:\s+(?:pro\s*forma|proforma|d.avoir|avoir|simplifi[ée]e|de\s+doit|commerciale))?\s*"
               r"(?:n\s*[°ºo˚]|n[uú]m[ée]ro|no\.?|#|:)\s*[:.\-]?\s*" + _CODE),
    re.compile(r"(?i)\binvoice\s*(?:no\.?|number|num|#|:)\s*[:.\-]?\s*" + _CODE),
    re.compile(r"(?i)\b(?:n\s*[°ºo˚]\s*(?:de\s*)?facture|num[ée]ro\s*(?:de\s*)?facture)\s*[:.\-]?\s*" + _CODE),
    re.compile(r"(?:فاتورة\s*(?:رقم|عدد|رق)|رقم\s*الفاتورة|عدد\s*الفاتورة)\s*[:\-.]?\s*" + _CODE),
]
_INV_HEAD = re.compile(r"(?i)\bfacture\b|\binvoice\b|فاتورة")
_INV_NEXT = re.compile(r"(?i)^\s*(?:n\s*[°ºo˚]|no\.?|رقم|#)\s*[:.\-]?\s*" + _CODE)

_ID_LABELS = {
    "nif": r"(?:\bN\.?\s?I\.?\s?F\b|identification\s+fiscale|التعريف\s*الجبائي|الرقم\s*الجبائي)",
    "nis": r"(?:\bN\.?\s?I\.?\s?S\b|identification\s+statistique|التعريف\s*الإحصائي|التعريف\s*الاحصائي|الرقم\s*الإحصائي)",
    "rc": r"(?:\bR\.?\s?C\b\.?|registre\s*(?:de\s*)?commerce|السجل\s*التجاري|رقم\s*السجل)",
    "ai": r"(?:\bA\.?\s?[Il1]\b\.?|article\s*(?:d.?\s*)?imposition|مادة\s*الإخضاع|مادة\s*الاخضاع|رقم\s*المادة)",
}
_ID_SEP = r"\s*(?:n\s*[°ºo˚]|رقم|num[ée]ro|no\.?)?\s*[:.\-]*\s*"
_ID_VALUE = {
    "nif": r"(\d[\d .\-]{8,28}\d)",
    "nis": r"(\d[\d .\-]{8,28}\d)",
    "ai": r"(\d[\d .\-]{5,20}\d)",
    "rc": r"(\d[0-9A-Za-z\u0621-\u064A/\- ]{4,28})",
}
_STOP_WORDS = re.compile(r"(?i)\b(?:NIF|NIS|A[Il1]|RC|N°|TVA|TEL|T[ée]l|FAX|EMAIL|E-?MAIL|RIB|CCP|BP|Adresse)\b|@")

_CLIENT_LABEL = re.compile(r"(?i)\b(?:client|doit|destinataire|adress[ée]e?\s*[àa]|factur[ée]e?\s*[àa]|bill\s*to|customer|sold\s*to)\b|"
                           r"الزبون|العميل|السيد|السادة|المشتري")
_COMPANY = re.compile(r"(?i)\b(?:SARL|EURL|SPA|S\.P\.A|SNC|SCS|ETS|[ÉE]tablissements?|Entreprise|STE|Soci[ée]t[ée]|Groupe|Cabinet|Pharmacie)\b|"
                      r"شركة|مؤسسة|مجمع|ذات\s*المسؤولية|م\.م\.ذ|ش\.ذ\.م\.م|صيدلية|مكتب")
_SKIP_SUPPLIER = re.compile(r"(?i)^(?:facture|invoice|bon\s|devis|proforma|duplicata|original|فاتورة|وصل)\b")

_PAY = [
    (re.compile(r"(?i)esp[èe]ces?|\bcash\b|نقد"), "نقدًا"),
    (re.compile(r"(?i)ch[èe]que|\bcheck\b|شيك|صك"), "شيك"),
    (re.compile(r"(?i)virement|transfer|تحويل"), "تحويل بنكي"),
    (re.compile(r"(?i)\bcarte\b|\bcard\b|بطاقة"), "بطاقة"),
    (re.compile(r"(?i)cr[ée]dit|à\s*terme|آجل|اجل"), "آجل"),
]
_CURRENCY = re.compile(r"(?i)\b(?:DZD|DA|DINARS?)\b|دج|د\.ج|دينار")

# تسميات المبالغ. الترتيب مهم: الأخص أولًا.
_AMT_LABELS = [
    ("timbre", re.compile(r"(?i)timbre|حقوق\s*الطابع|الطابع|طابع")),
    ("remise", re.compile(r"(?i)\bremise\b|rabais|discount|تخفيض|حسم")),
    ("net", re.compile(r"(?i)net\s*[àa]\s*payer|total\s*[àa]\s*payer|[àa]\s*payer|amount\s*due|صافي|"
                       r"المبلغ\s*(?:المستحق|الواجب\s*دفعه)|الواجب\s*دفعه")),
    ("ttc", re.compile(r"(?i)\bT\.?\s?T\.?\s?C\b\.?|toutes\s*taxes|total\s*g[ée]n[ée]ral|grand\s*total|"
                       r"المجموع\s*(?:مع|بالرسم|بكل|الإجمالي|الاجمالي|العام|الكلي)|المبلغ\s*(?:الإجمالي|الاجمالي|الكلي|مع\s*الرسم)|"
                       r"شامل\s*الرسم|بالرسوم|مع\s*الرسم")),
    ("ht", re.compile(r"(?i)\bH\.?\s?T\b\.?|hors\s*tax|sous[\s-]*total|المجموع\s*(?:خارج|بدون|قبل)|خارج\s*الرسم|"
                      r"بدون\s*رسم|قبل\s*الرسم|المبلغ\s*خارج|\bnet\s*commercial|subtotal")),
    ("tva", re.compile(r"(?i)\bT\.?\s?V\.?\s?A\b\.?|\bVAT\b|الرسم\s*على\s*القيمة|ر\.\s?ق\.\s?م|رسم\s*القيمة")),
    ("total", re.compile(r"(?i)\btotal\b|المجموع|الإجمالي|الاجمالي|\bmontant\b|المبلغ")),
]
_TOTAL_STOP = re.compile(r"(?i)\b(?:total|montant|sous|net|timbre|remise|tva|vat|t\.?t\.?c|h\.?t)\b|المجموع|الإجمالي|الاجمالي|الرسم|الطابع|صافي|المبلغ|خارج")
_HEADER_WORDS = re.compile(r"(?i)d[ée]signation|description|libell[ée]|article|produit|البيان|الوصف|التعيين|المنتج|السلعة")
_HEADER_QTY = re.compile(r"(?i)qt[ée]|quantit|qty|prix|\bp\.?u\b|unitaire|الكمية|السعر|الوحدة|المبلغ")
_NOT_ITEM = re.compile(r"(?i)\b(?:tel|t[ée]l|fax|nif|nis|rc|ai|rib|ccp|facture|invoice|date|bp)\b|@|www\.|فاتورة|هاتف|فاكس")

KEY_FIELDS = ["invoice_number", "date", "supplier", "ht", "tva", "ttc"]


def _clean_value(s):
    s = _STOP_WORDS.split(s, maxsplit=1)[0] if s else s
    return re.sub(r"\s+", " ", (s or "")).strip(" :-.,;|_")


def _digits_only(s):
    return re.sub(r"\D", "", s or "")


# ------------------------------------------------------------- استخراج الحقول
def _find_invoice_number(lines):
    for i, ln in enumerate(lines):
        for pat in _INV_PATTERNS:
            m = pat.search(ln)
            if m:
                v = m.group(1).strip("-_./")
                if any(c.isdigit() for c in v) and len(v) >= 2:
                    return v, ln
    for i, ln in enumerate(lines):
        if _INV_HEAD.search(ln):
            for nxt in lines[i + 1:i + 4]:
                m = _INV_NEXT.match(nxt)
                if m and any(c.isdigit() for c in m.group(1)):
                    return m.group(1).strip("-_./"), nxt
    return None, None


def _find_dates(lines):
    best, due = None, None
    for i, ln in enumerate(lines):
        ds = _dates_in(ln)
        if not ds:
            continue
        is_due = bool(_DUE_WORDS.search(ln))
        score = (3 if _DATE_WORDS.search(ln) else 0) + (1 if i < max(6, len(lines) // 3) else 0)
        for d in ds:
            if is_due:
                due = due or (d, ln)
            elif best is None or score > best[0]:
                best = (score, d, ln)
    return (best[1], best[2]) if best else (None, None), (due or (None, None))


def _find_ids(lines, client_line):
    """يعيد قوائم (قيمة, رقم_السطر) لكل معرّف، ثم يوزّعها بين المورّد والزبون."""
    found = {k: [] for k in _ID_LABELS}
    for i, ln in enumerate(lines):
        for key, lab in _ID_LABELS.items():
            for m in re.finditer(lab + _ID_SEP + _ID_VALUE[key], ln, flags=re.I):
                raw = m.group(1)
                if key == "rc":
                    raw = _clean_value(raw)
                    val = re.sub(r"\s*([/\-])\s*", r"\1", raw)
                    if len(_digits_only(val)) < 5:
                        continue
                else:
                    val = _digits_only(raw)
                found[key].append((val, i, ln))
    sup, cli = {}, {}
    for key, items in found.items():
        for val, i, ln in items:
            in_client = client_line is not None and client_line <= i <= client_line + 8
            tgt = cli if in_client else sup
            tgt.setdefault(key, (val, ln))
    return sup, cli


def _find_supplier(lines):
    head = []
    for l in lines[:14]:
        if _CLIENT_LABEL.search(l):
            break                                  # بعد هذا السطر يبدأ قسم الزبون
        if l:
            head.append(l)
    for ln in head:
        if _COMPANY.search(ln) and not _SKIP_SUPPLIER.match(ln):
            return _clean_value(re.sub(r"^[^\w\u0600-\u06FF]+", "", ln)), ln
    for ln in head:
        letters = len(re.findall(r"[A-Za-z\u0600-\u06FF]", ln))
        if letters >= 4 and not _SKIP_SUPPLIER.match(ln) and not _dates_in(ln) and not _NOT_ITEM.search(ln):
            return _clean_value(re.sub(r"^[^\w\u0600-\u06FF]+", "", ln)), ln
    return None, None


def _find_client(lines):
    for i, ln in enumerate(lines):
        m = _CLIENT_LABEL.search(ln)
        if not m:
            continue
        rest = _clean_value(ln[m.end():])
        if len(re.findall(r"[A-Za-z\u0600-\u06FF]", rest)) >= 3:
            return rest, ln, i
        for nxt in lines[i + 1:i + 3]:
            if len(re.findall(r"[A-Za-z\u0600-\u06FF]", nxt)) >= 3 and not _NOT_ITEM.search(nxt):
                return _clean_value(nxt), nxt, i
        return None, ln, i
    return None, None, None


def _labels_in(line):
    """كل التسميات المالية في السطر مع مواضعها. التسمية العامة 'total' تُهمل إن وُجدت تسمية أدق."""
    hits = []
    for key, pat in _AMT_LABELS:
        for m in pat.finditer(line):
            if any(m.start() < e and s < m.end() for _, s, e in hits):
                continue
            hits.append((key, m.start(), m.end()))
    specific = [h for h in hits if h[0] != "total"]
    return sorted(specific if specific else hits, key=lambda h: h[1])


def _find_amounts(lines):
    """يمرّ على الأسطر ويسند كل مبلغ لأقرب تسمية قبله (أو بعده إن لم يسبقه شيء)."""
    got = {}      # field -> list[(value, line, rate)]
    pending = None
    for i, ln in enumerate(lines):
        labs = _labels_in(ln)
        amts = _amounts_in(ln)
        cands = [parse_amount(x) for x in re.findall(r"(\d{1,2}(?:[.,]\d)?)\s*%|%\s*(\d{1,2}(?:[.,]\d)?)", ln) for x in x if x]
        rate = next((c for c in cands if c in (19.0, 9.0, 7.0, 0.0)), cands[-1] if cands else None)
        if pending and amts and not labs and not re.search(r"[A-Za-z\u0600-\u06FF]", ln):
            got.setdefault(pending, []).append((amts[0][1], ln, rate))
            pending = None
            continue
        pending = None
        if not labs:
            continue
        if not amts:
            # لا مبلغ في سطر التسمية: قد يكون المبلغ في السطر التالي، أو هي ترويسة أعمدة (HT | TVA | TTC)
            nxt = _amounts_in(lines[i + 1]) if i + 1 < len(lines) else []
            real = [l for l in labs if l[0] != "total"]
            if len(real) >= 2 and len(nxt) >= len(real) and not _labels_in(lines[i + 1]):
                for (key, _, _), (_, v, _) in zip(real, nxt):
                    got.setdefault(key, []).append((v, lines[i + 1], rate))
            elif labs:
                pending = labs[-1][0]
            continue
        used = set()
        for k, (key, s, e) in enumerate(labs):
            nxt_s = labs[k + 1][1] if k + 1 < len(labs) else 10 ** 6
            mine = [a for a in amts if s <= a[0] < nxt_s and a[0] not in used]
            if not mine and len(labs) == 1:
                mine = [a for a in amts if a[0] < s]          # المبلغ قبل التسمية (ترتيب معكوس)
            if mine:
                a = mine[-1]
                used.add(a[0])
                got.setdefault(key, []).append((a[1], ln, rate))
    return got


def _pick(entries, prefer="last"):
    if not entries:
        return None, None
    e = entries[-1] if prefer == "last" else entries[0]
    return e[0], e[1]


def _find_items(lines, first_total_idx):
    start = 0
    for i, ln in enumerate(lines):
        if _HEADER_WORDS.search(ln) and _HEADER_QTY.search(ln):
            start = i + 1
            break
    end = first_total_idx if first_total_idx is not None else len(lines)
    header_found = start > 0
    items = []
    for ln in lines[start:end]:
        if not ln or _NOT_ITEM.search(ln) or _TOTAL_STOP.search(ln) and not header_found:
            continue
        it = _parse_item_line(ln, header_found)
        if it:
            items.append(it)
    # أرقام الأسطر التسلسلية (1،2،3…) ملتصقة بالوصف: تُحذف فقط إن كانت متسلسلة فعلًا
    lead = [re.match(r"^(\d{1,3})\s+(?=[A-Za-z\u00C0-\u00FF\u0600-\u06FF])", i["desc"]) for i in items]
    if items and all(lead) and [int(m.group(1)) for m in lead] == list(range(1, len(items) + 1)):
        for i, m in zip(items, lead):
            i["desc"] = i["desc"][m.end():]
    return items


def _parse_item_line(line, header_found):
    clean = _PCT.sub(" ", line)
    clean = re.sub(r"^\s*\d{1,3}\s*[.\-)]\s+(?=\D)", "", clean)          # رقم السطر "1." أو "1-"
    toks = list(re.finditer(r"\S+", clean))
    # الكتلة الرقمية في آخر السطر
    j = len(toks)
    while j > 0 and re.fullmatch(r"[-+]?\d[\d.,]*", toks[j - 1].group(0)):
        j -= 1
    num_toks = [t.group(0) for t in toks[j:]]
    desc = clean[:toks[j].start()].strip() if j < len(toks) else ""
    if len(num_toks) < 2 or len(num_toks) > 9:
        return None
    if len(re.findall(r"[A-Za-z\u0600-\u06FF]", desc)) < 2:
        return None
    # نجرّب كل تقسيمات الأرقام المتجاورة (لأن "5 190,00" قد تكون 5 و 190,00 أو 5190,00)
    best = None
    n = len(num_toks)
    for mask in itertools.product((0, 1), repeat=n - 1):
        groups, cur = [], [num_toks[0]]
        for tok, merge in zip(num_toks[1:], mask):
            if merge:
                cur.append(tok)
            else:
                groups.append(cur)
                cur = [tok]
        groups.append(cur)
        vals = [parse_amount(" ".join(g)) for g in groups]
        if any(v is None for v in vals) or len(vals) < 2:
            continue
        k = len(vals) - 1                               # المجموع هو آخر قيمة في السطر
        total = vals[k]
        if total <= 0:
            continue
        if k >= 2:
            qty, pu = vals[k - 2], vals[k - 1]
            if qty > 0 and pu > 0 and abs(qty * pu - total) <= max(0.05, 0.005 * abs(total)):
                cand = (3, qty, pu, total)
                if not best or cand[3] > best[3]:
                    best = cand
        elif header_found and len(vals) == 2 and vals[0] > 0 and not best:
            best = (1, vals[0], round(total / vals[0], 4), total)
    if not best:
        return None
    _, qty, pu, total = best
    return {"desc": re.sub(r"\s+", " ", desc).strip(" -|:"), "qty": qty, "unit_price": pu, "total": total,
            "ok": best[0] == 3}


# ------------------------------------------------------------------ الواجهة
def extract_invoice(text):
    """النقطة الرئيسية: نص OCR → {'fields','meta','items','warnings','score'}."""
    text = fix_arabic_order(normalize_text(text))
    if len(_AR.findall(text)) > 0.2 * max(1, len(re.findall(r"[A-Za-z]", text)) + len(_AR.findall(text))):
        text = heal_arabic(text)
    lines = [l for l in text.split("\n")]
    nonempty = [l for l in lines if l]
    fields, meta = {}, {}

    def put(key, value, line=None, status="ok", note=None):
        fields[key] = value
        meta[key] = {"status": status if value not in (None, "") else "missing", "line": line, "note": note}

    num, ln = _find_invoice_number(nonempty)
    put("invoice_number", num, ln)
    (d, dl), (due, duel) = _find_dates(nonempty)
    put("date", d.isoformat() if d else None, dl)
    put("due_date", due.isoformat() if due else None, duel)
    sup, supl = _find_supplier(nonempty)
    put("supplier", sup, supl, "ok" if sup else "missing")
    cli, clil, cli_idx = _find_client(nonempty)
    put("client", cli, clil)
    sids, cids = _find_ids(nonempty, cli_idx)
    shape = _ids_by_shape(nonempty, cli_idx)
    for key in ("nif", "nis", "rc", "ai"):
        v = sids.get(key)
        if not v and shape.get(key):
            v = shape[key]
            put("supplier_" + key, v[0], v[1], "warn", "صُنِّف بشكل الرقم لا بتسميته — راجعه")
        else:
            put("supplier_" + key, v[0] if v else None, v[1] if v else None)
        if key == "rc" and fields.get("supplier_rc"):
            fixed, ch = _fix_rc(fields["supplier_rc"])
            if ch:
                fields["supplier_rc"] = fixed
                meta["supplier_rc"]["status"], meta["supplier_rc"]["note"] = "warn", "صُحِّح حرف RC (8→B) — راجعه"
        v = cids.get(key)
        put("client_" + key, v[0] if v else None, v[1] if v else None)

    pay = next((name for pat, name in _PAY if pat.search(text)), None)
    put("payment_method", pay)
    put("currency", "DZD" if _CURRENCY.search(text) else None, None, "ok")

    amts = _find_amounts(nonempty)
    first_total_idx = None
    for i, l in enumerate(nonempty):
        if any(k in ("ht", "tva", "ttc", "net", "timbre") for k, _, _ in _labels_in(l)) and _amounts_in(l):
            first_total_idx = i
            break
    for key in ("ht", "ttc", "net", "timbre", "remise"):
        v, l = _pick(amts.get(key))
        put(key, v, l)
    # TVA: قد تكون عدة أسطر بنسب مختلفة
    tva_entries = amts.get("tva", [])
    tva_total, tva_line, rates = None, None, []
    if tva_entries:
        rates = sorted({e[2] for e in tva_entries if e[2]}, reverse=True)
        distinct = {(round(e[0], 2), e[2]) for e in tva_entries}
        if len(distinct) >= 2 and len(rates) >= 2:
            tva_total = round(sum(v for v, _ in distinct), 2)
            tva_line = "؛ ".join(e[1] for e in tva_entries)
        else:
            tva_total, tva_line = tva_entries[-1][0], tva_entries[-1][1]
    put("tva", tva_total, tva_line)
    if not rates:
        m = re.search(r"(?i)(?:tva|vat|الرسم)[^\n%]{0,30}?(\d{1,2}(?:[.,]\d)?)\s*%", text)
        if m:
            rates = [parse_amount(m.group(1))]
    put("tva_rate", rates[0] if rates else None, None)
    fields["tva_rates"] = rates
    put("items_note", None)

    # مبالغ مشتقّة عند غياب الحقل (تُوسَم computed ولا تُعتبر قراءة)
    def compute(key, value, note):
        if fields.get(key) is None and value is not None:
            fields[key] = round(value, 2)
            meta[key] = {"status": "computed", "line": None, "note": note}

    ht, tva, ttc, net, timbre = (fields.get(k) for k in ("ht", "tva", "ttc", "net", "timbre"))
    if ttc is None and net is not None:
        compute("ttc", net - (timbre or 0), "محسوب: الصافي − الطابع")
    ttc = fields.get("ttc")
    if ttc is None and ht is not None and tva is not None:
        compute("ttc", ht + tva, "محسوب: HT + TVA")
    ttc = fields.get("ttc")
    if tva is None and ht is not None and ttc is not None:
        compute("tva", ttc - ht, "محسوب: TTC − HT")
    if ht is None and ttc is not None and tva is not None:
        compute("ht", ttc - tva, "محسوب: TTC − TVA")
    if fields.get("ttc") is None:
        g = amts.get("total")
        if g:
            put("ttc", g[-1][0], g[-1][1], "warn", "من تسمية عامة (مجموع) — راجعه")

    items = _find_items(nonempty, first_total_idx)
    fields["items_count"] = len(items)
    res = {"fields": fields, "meta": meta, "items": items}
    res["warnings"] = validate(res)
    res["score"] = score(res)
    return res


# ------------------------------------------------------------------ التحقق
def _close(a, b, rel=0.0, absol=0.11):
    return abs(a - b) <= max(absol, rel * max(abs(a), abs(b)))


def _is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def validate(res):
    """فحوص حسابية ومنطقية. تعيد قائمة (مستوى, حقل, رسالة). المستوى: error | warn | info."""
    f, meta, items = res["fields"], res["meta"], res.get("items") or []
    W = []
    for key in ("ht", "tva", "ttc", "net", "timbre", "remise"):
        value = f.get(key)
        if value in (None, ""):
            continue
        if not _is_number(value):
            W.append(("error", key, "قيمة غير رقمية — صحّح الحقل."))
            f[key] = None
    rates = [r for r in (f.get("tva_rates") or []) if _is_number(r)]
    f["tva_rates"] = rates
    ht, tva, ttc, net, timbre = (f.get(k) for k in ("ht", "tva", "ttc", "net", "timbre"))
    rem = f.get("remise") or 0

    for k, label in (("invoice_number", "رقم الفاتورة"), ("date", "التاريخ"), ("supplier", "اسم المورّد")):
        if not f.get(k):
            W.append(("warn", k, f"لم يُعثر على {label}."))
    if ttc is None and ht is None:
        W.append(("error", "ttc", "لم يُعثر على أي مبلغ إجمالي — راجع النص يدويًا."))

    if ht is not None and tva is not None and ttc is not None and meta.get("ttc", {}).get("status") != "computed" \
            and meta.get("tva", {}).get("status") != "computed" and meta.get("ht", {}).get("status") != "computed":
        if not _close(ht + tva, ttc):
            W.append(("error", "ttc", f"HT + TVA = {ht + tva:,.2f} لكن TTC المقروء = {ttc:,.2f}. أحد الأرقام مقروء خطأً."))
    if ttc is not None and net is not None and timbre is not None and not _close(ttc + timbre, net):
        W.append(("error", "net", f"TTC + الطابع = {ttc + timbre:,.2f} لكن الصافي المقروء = {net:,.2f}."))
    rates = f.get("tva_rates") or []
    if ht is not None and tva is not None and len(rates) == 1:
        base = ht - rem if rem and not _close(ht, ttc or 0) else ht
        exp = base * rates[0] / 100.0
        if not _close(exp, tva, rel=0.0002, absol=1.0):
            W.append(("warn", "tva", f"TVA المتوقعة بنسبة {rates[0]:g}% هي {exp:,.2f} لكن المقروء {tva:,.2f}."))
    if rates and any(r not in (0, 9, 19) for r in rates):
        W.append(("info", "tva_rate", f"نسبة TVA {', '.join(f'{r:g}%' for r in rates)} غير معتادة — تحقق منها."))
    if items:
        s = round(sum(i["total"] for i in items), 2)
        if ht is not None and _close(s, ht, rel=0.001, absol=0.5):
            pass
        elif ttc is not None and _close(s, ttc, rel=0.001, absol=0.5):
            W.append(("info", "items", "مجموع البنود يطابق TTC، أي أن أسعار البنود تشمل الرسم."))
        elif ht is not None:
            W.append(("warn", "items", f"مجموع البنود المقروءة {s:,.2f} لا يطابق HT = {ht:,.2f} (قد تكون بنود لم تُقرأ)."))
        bad = [i for i in items if not i["ok"]]
        if bad:
            W.append(("warn", "items", f"{len(bad)} بند لم تتطابق فيه الكمية × السعر مع المجموع."))

    def idlen(key, label, ok_lens):
        v = f.get(key)
        if v and len(_digits_only(v)) not in ok_lens:
            W.append(("warn", key, f"{label} طوله {len(_digits_only(v))} خانة — خلاف الطول المعتاد، راجعه على الأصل."))
    inv = f.get("invoice_number") or ""
    if re.search(r"(?<=\d)[OolI|]|[OolI|](?=\d)", inv):
        W.append(("warn", "invoice_number", "رقم الفاتورة يحوي حرفًا قد يكون رقمًا (O/0 أو l/1) — قارنه بالأصل."))
    idlen("supplier_nif", "NIF المورّد", (15, 20))
    idlen("supplier_nis", "NIS المورّد", (15,))
    idlen("client_nif", "NIF الزبون", (15, 20))

    if f.get("date"):
        try:
            d = _dt.date.fromisoformat(f["date"])
            if d > _dt.date.today() + _dt.timedelta(days=1):
                W.append(("warn", "date", "تاريخ الفاتورة في المستقبل — غالبًا خطأ قراءة."))
            elif d.year < 2000:
                W.append(("warn", "date", "تاريخ قديم جدًا — غالبًا خطأ قراءة."))
        except ValueError:
            pass
    for k, m in meta.items():
        if m["status"] == "computed":
            W.append(("info", k, f"{k}: {m['note']}"))
    return W


def score(res):
    """نسبة ثقة تقريبية 0-100: حضور الحقول الأساسية ناقص عقوبات الأخطاء الحسابية."""
    f, meta = res["fields"], res["meta"]
    got = sum(1 for k in KEY_FIELDS if f.get(k) not in (None, "") and meta[k]["status"] in ("ok", "computed"))
    s = 100.0 * got / len(KEY_FIELDS)
    for lvl, _, _ in res["warnings"]:
        s -= 18 if lvl == "error" else 6 if lvl == "warn" else 0
    return int(max(0, min(100, round(s))))


def field_status(res, key):
    """حالة عرض الحقل في الواجهة: ok / warn / error / missing / computed."""
    m = res["meta"].get(key, {})
    st = m.get("status", "missing")
    for lvl, k, _ in res["warnings"]:
        if k == key and lvl in ("error", "warn"):
            return "error" if lvl == "error" else "warn"
    return st


# ------------------------------------------------------------- OCR → استخراج
def _prepare_for_ocr(im):
    """رمادي + تكبير إن كانت الصورة صغيرة (Tesseract يحتاج ~300dpi ≈ عرض 2400px لـ A4)."""
    from PIL import Image
    g = im.convert("L")
    if g.width < 1900:
        k = 2400.0 / g.width
        g = g.resize((int(g.width * k), int(g.height * k)), Image.Resampling.LANCZOS)
    return g


def _merge(primary, secondary):
    """يملأ حقول primary الناقصة من secondary ثم يعيد التحقق."""
    changed = False
    for k, v in secondary["fields"].items():
        if k in ("items_count", "tva_rates"):
            continue
        if primary["fields"].get(k) in (None, "") and v not in (None, ""):
            primary["fields"][k] = v
            m = dict(secondary["meta"].get(k, {}))
            m["status"] = "warn" if m.get("status") == "ok" else m.get("status", "warn")
            m["note"] = "من قراءة ثانية — راجعه"
            primary["meta"][k] = m
            changed = True
    if not primary["items"] and secondary["items"]:
        primary["items"] = secondary["items"]
        changed = True
    if changed:
        primary["warnings"] = validate(primary)
        primary["score"] = score(primary)
    return primary


def _ocr_lines(pytesseract, g, lang, config, psm):
    """أسطر OCR مع موضعها الرأسي: [(نص, y_mid, ارتفاع)]."""
    from pytesseract import Output
    d = pytesseract.image_to_data(g, lang=lang, config=f"{config} --psm {psm}".strip(), output_type=Output.DICT)
    groups = {}
    for i, t in enumerate(d["text"]):
        t = (t or "").strip()
        if not t or float(d["conf"][i]) < 0:
            continue
        groups.setdefault((d["block_num"][i], d["par_num"][i], d["line_num"][i]), []).append(
            (d["word_num"][i], t, d["top"][i], d["height"][i]))
    lines = []
    for ws in groups.values():
        ws.sort()
        top = min(w[2] for w in ws)
        bot = max(w[2] + w[3] for w in ws)
        lines.append((" ".join(w[1] for w in ws), (top + bot) / 2.0, bot - top))
    lines.sort(key=lambda x: x[1])
    return lines


def _align_numbers(ara_lines, num_lines):
    """مستند عربي: التسميات من قراءة العربية (بعد حذف أرقامها غير الموثوقة) + الأرقام من قراءة الإنجليزية
    في نفس الارتفاع. قراءة الأرقام بـ eng أدق بكثير من ara."""
    hs = sorted(h for _, _, h in ara_lines) or [30]
    tol = max(10.0, 0.7 * hs[len(hs) // 2])
    labels = fix_arabic_order("\n".join(" ".join(t for t in txt.split() if not re.search(r"\d", t))
                                         for txt, _, _ in ara_lines))
    out, used = [], set()
    for (txt, y, _), lab in zip(ara_lines, labels.split("\n")):
        nums = []
        for j, (t2, y2, _) in enumerate(num_lines):
            if abs(y2 - y) <= tol:
                used.add(j)
                nums += [t.strip(".,:;") for t in normalize_text(t2).split() if re.search(r"\d", t)]
        out.append((lab + " " + " ".join(nums)).strip())
    for j, (t2, _, _) in enumerate(num_lines):               # أسطر أرقام بلا مقابل عربي
        if j not in used:
            nums = [t.strip(".,:;") for t in normalize_text(t2).split() if re.search(r"\d", t)]
            if nums:
                out.append(" ".join(nums))
    return "\n".join(out)


def _arabic_ratio(text):
    ar = len(_AR.findall(text))
    la = len(re.findall(r"[A-Za-z]", text))
    return ar / max(1, ar + la)


def _good(r):
    return r["score"] >= 90 and not any(l == "error" for l, _, _ in r["warnings"])


def extract_from_image(im, pytesseract, lang, config="", progress=None):
    """OCR تكيّفي: قراءة سريعة أولًا، وإن لم تكفِ قراءات إضافية ثم دمج الحقول الناقصة.
    1) كل اللغات psm 6. 2) إن لزم psm 4. 3) إن كان المستند عربيًا: عربي فقط + قراءة أرقام (eng) محاذاة بالموضع.
    4) إن كان فرنسيًا: fra+eng فقط (يزيل هلوسة الحروف العربية). progress(msg) اختياري."""
    say = progress or (lambda m: None)
    g = _prepare_for_ocr(im)
    say("قراءة النص (1)…")
    txt = pytesseract.image_to_string(g, lang=lang, config=f"{config} --psm 6".strip())
    best = extract_invoice(txt)
    best["text"], best["psm"] = best_text(txt), 6
    if _good(best):
        return best
    say("قراءة ثانية بنمط تقسيم مختلف…")
    try:
        t2 = pytesseract.image_to_string(g, lang=lang, config=f"{config} --psm 4".strip())
        r2 = extract_invoice(t2)
        r2["text"], r2["psm"] = best_text(t2), 4
        if (r2["score"], _filled(r2)) > (best["score"], _filled(best)):
            best, r2 = r2, best
        best = _merge(best, r2)
    except Exception:
        pass
    if _good(best):
        return best
    langs = set(lang.split("+"))
    try:
        if _arabic_ratio(best["text"]) > 0.25 and "ara" in langs:
            say("مستند عربي: قراءة العربية والأرقام منفصلتين…")
            ara = _ocr_lines(pytesseract, g, "ara", config, 6)
            eng_l = "eng" if "eng" in langs else lang
            num = _ocr_lines(pytesseract, g, eng_l, config, 6)
            t3 = _align_numbers(ara, num)
            r3 = extract_invoice(t3)
            r3["text"], r3["psm"] = best_text(t3), "ara+num"
            if (r3["score"], _filled(r3)) > (best["score"], _filled(best)):
                best, r3 = r3, best
            best = _merge(best, r3)
        elif langs & {"fra", "eng"}:
            say("قراءة بالفرنسية والإنجليزية فقط…")
            l2 = "+".join(x for x in ("fra", "eng") if x in langs)
            t4 = pytesseract.image_to_string(g, lang=l2, config=f"{config} --psm 6".strip())
            r4 = extract_invoice(t4)
            r4["text"], r4["psm"] = best_text(t4), "fra+eng"
            if (r4["score"], _filled(r4)) > (best["score"], _filled(best)):
                best, r4 = r4, best
            best = _merge(best, r4)
    except Exception:
        pass
    return best


def best_text(t):
    return fix_arabic_order(normalize_text(t))


def _filled(r):
    return sum(1 for v in r["fields"].values() if v not in (None, "", []))


# --------------------------------------------------------------------- التصدير
COLUMNS = [
    ("source", "المصدر"), ("invoice_number", "رقم الفاتورة"), ("date", "التاريخ"), ("due_date", "الاستحقاق"),
    ("supplier", "المورّد"), ("supplier_nif", "NIF المورّد"), ("supplier_nis", "NIS المورّد"),
    ("supplier_rc", "RC المورّد"), ("supplier_ai", "AI المورّد"),
    ("client", "الزبون"), ("client_nif", "NIF الزبون"), ("client_nis", "NIS الزبون"),
    ("client_rc", "RC الزبون"), ("client_ai", "AI الزبون"),
    ("ht", "المجموع خارج الرسم HT"), ("tva_rate", "نسبة TVA %"), ("tva", "مبلغ TVA"), ("timbre", "حقوق الطابع"),
    ("ttc", "المجموع TTC"), ("net", "الصافي للدفع"), ("payment_method", "طريقة الدفع"),
    ("currency", "العملة"), ("score", "الثقة %"), ("warnings", "تنبيهات"),
]
ID_COLUMNS = {"supplier_nif", "supplier_nis", "supplier_rc", "supplier_ai",
              "client_nif", "client_nis", "client_rc", "client_ai", "invoice_number"}
MONEY_COLUMNS = {"ht", "tva", "timbre", "ttc", "net"}


def to_row(res, source=""):
    f = dict(res["fields"])
    f["source"] = source or res.get("source", "")
    f["score"] = res.get("score")
    f["warnings"] = " | ".join(m for lvl, _, m in res.get("warnings", []) if lvl in ("error", "warn"))
    return {k: f.get(k) for k, _ in COLUMNS}


def _safe_cell(value):
    """Prevent CSV/Excel formula injection from text read by OCR."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _csv_cell(key, value):
    if value is None:
        return ""
    if key in ID_COLUMNS and value != "":
        # CSV يبقى صالحًا لأي مستورد (ERP/pandas): أرقام صافية بلا ="..."،
        # وملف Excel هو الذي يتولى تثبيت الأصفار البادئة كنص.
        value = str(value)
    return _safe_cell(value)


def write_csv(path, results):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow([h for _, h in COLUMNS])
        for r in results:
            row = to_row(r)
            w.writerow([_csv_cell(k, row[k]) for k, _ in COLUMNS])


def write_json(path, results):
    out = []
    for r in results:
        out.append({"source": r.get("source", ""), "fields": r["fields"], "items": r["items"],
                    "warnings": [{"level": a, "field": b, "message": c} for a, b, c in r["warnings"]],
                    "score": r["score"]})
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=2)


def write_xlsx(path, results):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "الفواتير"
    ws.sheet_view.rightToLeft = True
    ws.append([h for _, h in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in results:
        row = to_row(r)
        ws.append([row[k] for k, _ in COLUMNS])
        for idx, (k, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=ws.max_row, column=idx)
            if k in ID_COLUMNS and cell.value not in (None, ""):
                cell.value = str(cell.value)
                cell.number_format = "@"                      # يحفظ الأصفار البادئة
            elif k in MONEY_COLUMNS and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"                          # لا تُفسَّر كنص OCR كصيغة
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=(k == "warnings"))
    for idx, (k, h) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(idx)].width = 48 if k == "warnings" else max(13, min(30, len(h) + 4))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("البنود")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["المصدر", "رقم الفاتورة", "الوصف", "الكمية", "سعر الوحدة", "المجموع", "تطابق الحساب"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
    for r in results:
        for it in r["items"]:
            ws2.append([r.get("source", ""), r["fields"].get("invoice_number"), it["desc"], it["qty"],
                        it["unit_price"], it["total"], "نعم" if it["ok"] else "راجعه"])
    for col, wd in zip("ABCDEFG", (18, 16, 46, 10, 14, 14, 14)):
        ws2.column_dimensions[col].width = wd
    for row in ws2.iter_rows(min_row=2):
        for c in row[4:6]:
            c.number_format = "#,##0.00"
    wb.save(path)


def erp_payload(res, source=""):
    """حمولة JSON عامة لإرسالها إلى نظام ERP. الحقول بأسماء إنجليزية ثابتة والمبالغ أرقام."""
    f = res["fields"]
    return {
        "source": "canon-scan-pro", "document": source,
        "invoice": {k: f.get(k) for k in (
            "invoice_number", "date", "due_date", "supplier", "supplier_nif", "supplier_nis", "supplier_rc",
            "supplier_ai", "client", "client_nif", "client_nis", "client_rc", "client_ai", "ht", "tva_rate", "tva",
            "timbre", "ttc", "net", "payment_method", "currency")},
        "items": [{"description": i["desc"], "quantity": i["qty"], "unit_price": i["unit_price"], "total": i["total"]}
                  for i in res["items"]],
        "needs_review": any(l in ("error", "warn") for l, _, _ in res["warnings"]),
        "confidence": res["score"],
    }
